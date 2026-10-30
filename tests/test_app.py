"""Backend tests using an in-memory fake Firestore. Run: python tests/test_app.py"""
import base64
import io
import os
import sys
from unittest.mock import MagicMock

import openpyxl

for m in ("firebase_admin", "firebase_admin.credentials", "firebase_admin.firestore"):
    sys.modules[m] = MagicMock()
os.environ["USERS"] = "juan:pw1,maria:pw2"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


class Snap:
    def __init__(s, ref, data):
        s.reference, s.id, s._d, s.exists = ref, ref.id, data, data is not None

    def to_dict(s):
        return dict(s._d) if s._d is not None else None


class Ref:
    def __init__(s, db, path):
        s.db, s.path, s.id = db, path, path[-1]

    def collection(s, n):
        return Col(s.db, s.path + (n,))

    def set(s, d):
        s.db.store[s.path] = dict(d)

    def get(s):
        return Snap(s, s.db.store.get(s.path))

    def delete(s):
        s.db.store.pop(s.path, None)


class Col:
    def __init__(s, db, path):
        s.db, s.path = db, path

    def document(s, i):
        return Ref(s.db, s.path + (i,))

    def _ids(s):
        return [p for p in s.db.store if len(p) == len(s.path) + 1 and p[:len(s.path)] == s.path]

    def stream(s):
        return [Snap(Ref(s.db, p), s.db.store[p]) for p in s._ids()]

    def list_documents(s):
        return [Ref(s.db, p) for p in s._ids()]


class Batch:
    def __init__(s):
        s.ops = []

    def set(s, r, d):
        s.ops.append(("s", r, d))

    def delete(s, r):
        s.ops.append(("d", r, None))

    def commit(s):
        assert len(s.ops) <= 500
        for k, r, d in s.ops:
            r.set(d) if k == "s" else r.delete()


class DB:
    def __init__(s):
        s.store = {}

    def collection(s, n):
        return Col(s, (n,))

    def batch(s):
        return Batch()


import app as A  # noqa: E402

A.db = DB()
A.db.store[("users", "juan", "students", "1")] = {"name": "Old Student"}  # pre-blocks data
A.db.store[("users", "juan", "attendance", "2026-10-01")] = {"1": "P"}

A.time.sleep = lambda s: None  # skip the wrong-password delay in tests


class Browser:
    """Pretends to be separate browsers: each (user, password) gets its own cookie jar and logs in."""

    def __init__(self):
        self.anon, self.jars = A.app.test_client(), {}

    def _client(self, headers):
        if not headers:
            return self.anon
        key = (headers["u"], headers["p"])
        if key not in self.jars:
            cl = A.app.test_client()
            ok_login = cl.post("/api/login", json={"username": key[0], "password": key[1]}).status_code == 200
            self.jars[key] = cl if ok_login else self.anon
        return self.jars[key]

    def get(self, *a, headers=None, **k): return self._client(headers).get(*a, **k)
    def post(self, *a, headers=None, **k): return self._client(headers).post(*a, **k)
    def put(self, *a, headers=None, **k): return self._client(headers).put(*a, **k)
    def patch(self, *a, headers=None, **k): return self._client(headers).patch(*a, **k)
    def delete(self, *a, headers=None, **k): return self._client(headers).delete(*a, **k)


c = Browser()
H = lambda u, p: {"u": u, "p": p}
J, M = H("juan", "pw1"), H("maria", "pw2")


def ok(label, cond):
    print(("PASS " if cond else "FAIL ") + label)
    assert cond, label


ok("no login -> 401", c.get("/api/blocks").status_code == 401)
ok("wrong password -> 401", c.get("/api/blocks", headers=H("juan", "x")).status_code == 401)
b = c.get("/api/blocks", headers=J).get_json()["blocks"]
ok("5 default blocks A-E", [x["name"] for x in b] == [f"Block {x}" for x in "ABCDE"])
d = c.get("/api/blocks/block-a/data", headers=J).get_json()
ok("older data moved into Block A", d["students"] == [{"id": "1", "name": "Old Student"}]
   and d["records"] == {"2026-10-01": {"1": "P"}} and d["activities"] == [])
ok("Block B starts empty", c.get("/api/blocks/block-b/data", headers=J).get_json()["students"] == [])

body = {"students": [{"id": "7", "name": "Maria B"}, {"id": "8", "name": "Jose B"}],
        "records": {"2026-10-14": {"7": "P", "8": "L", "99": "P"}}}
ok("PUT block B", c.put("/api/blocks/block-b/data", json=body, headers=J).status_code == 200)
d = c.get("/api/blocks/block-b/data", headers=J).get_json()
ok("Block B saved, unknown student ignored",
   len(d["students"]) == 2 and d["records"] == {"2026-10-14": {"7": "P", "8": "L"}})
ok("Block A untouched", len(c.get("/api/blocks/block-a/data", headers=J).get_json()["students"]) == 1)
ok("bad payload 400", c.put("/api/blocks/block-b/data", json={"students": "x"}, headers=J).status_code == 400)
ok("bad date 400", c.put("/api/blocks/block-b/data",
                         json={"students": [], "records": {"oct": {}}}, headers=J).status_code == 400)
ok("unknown block 404", c.get("/api/blocks/nope/data", headers=J).status_code == 404)
c.get("/api/blocks", headers=M)
ok("another login has separate data", c.get("/api/blocks/block-b/data", headers=M).get_json()["students"] == [])

r = c.post("/api/blocks", json={"name": "Block F"}, headers=J)
ok("add Block F", r.status_code == 200 and r.get_json()["id"] == "block-f")
ok("duplicate block name rejected", c.post("/api/blocks", json={"name": "block f"}, headers=J).status_code == 400)
ok("rename block", c.patch("/api/blocks/block-f", json={"name": "Block Z"}, headers=J).status_code == 200)

# ---- activities
body = {"students": [{"id": "1", "name": "Ana"}, {"id": "2", "name": "Ben"}], "records": {},
        "activities": [{"id": "a1", "title": "Quiz 1", "date": "2026-10-14", "done": ["1", "2", "99"]},
                       {"id": "a2", "title": "Poster", "date": "", "done": []}]}
ok("PUT with activities", c.put("/api/blocks/block-c/data", json=body, headers=J).status_code == 200)
d = c.get("/api/blocks/block-c/data", headers=J).get_json()
ok("activities saved, unknown student dropped",
   [(a["title"], a["done"]) for a in d["activities"]] == [("Quiz 1", ["1", "2"]), ("Poster", [])])
writes, orig = [], A.commit_ops
A.commit_ops = lambda ops: (writes.append(len(ops)), orig(ops))
c.put("/api/blocks/block-c/data", json=body, headers=J)
ok("saving unchanged data writes nothing", sum(writes) == 0)
body["activities"][0]["done"] = ["1"]
writes.clear()
c.put("/api/blocks/block-c/data", json=body, headers=J)
ok("one changed activity = one write", sum(writes) == 1)
A.commit_ops = orig
c.put("/api/blocks/block-c/data", json={"students": body["students"], "records": {}}, headers=J)
ok("page that sends no activities leaves them alone",
   len(c.get("/api/blocks/block-c/data", headers=J).get_json()["activities"]) == 2)
ok("bad activity id rejected", c.put("/api/blocks/block-c/data", headers=J,
   json=dict(body, activities=[{"id": "x y", "title": "Bad", "done": []}])).status_code == 400)
ok("empty activity name rejected", c.put("/api/blocks/block-c/data", headers=J,
   json=dict(body, activities=[{"id": "x", "title": "", "done": []}])).status_code == 400)

# ---- spreadsheets
r = c.get("/api/blocks/block-c/export.xlsx", headers=J)
wb = openpyxl.load_workbook(io.BytesIO(r.data))
ok("export: attendance + activities sheets", wb.sheetnames == ["Block C", "Block C Activities"])
rows = list(wb["Block C Activities"].iter_rows(min_row=1, max_row=3, values_only=True))
ok("activities header", rows[0][:2] == ("ID", "Name") and rows[0][4].startswith("Quiz 1")
   and rows[0][5].startswith("Poster"))
ok("done shown as Done", rows[1][4] == "Done" and rows[2][4] is None)
wb = openpyxl.load_workbook(io.BytesIO(c.get("/api/export-all.xlsx", headers=J).data))
ok("export all: sheet per block + activities",
   wb.sheetnames[:2] == ["Block A", "Block B"] and "Block C Activities" in wb.sheetnames)
r = c.post("/api/parse-import", headers=J, content_type="multipart/form-data",
           data={"file": (io.BytesIO(b"ID,Name,2026-10-16\n1,Ana,P\n"), "x.csv")})
ok("import spreadsheet", r.status_code == 200 and r.get_json()["records"] == {"2026-10-16": {"1": "P"}})

# ---- deleting
ok("delete block", c.delete("/api/blocks/block-c", headers=J).status_code == 200)
ok("its data and activities are gone", not [p for p in A.db.store if p[:4] == ("users", "juan", "blocks", "block-c")])
for x in ["block-a", "block-b", "block-d", "block-e", "block-f"]:
    c.delete(f"/api/blocks/{x}", headers=J)
ok("cannot delete the last block", c.delete("/api/blocks/block-z", headers=J).status_code in (400, 404)
   and len(c.get("/api/blocks", headers=J).get_json()["blocks"]) == 1)

# ---- login page and sessions
anon = A.app.test_client()
r = anon.get("/")
ok("app page redirects to the login page", r.status_code == 302 and r.headers["Location"].endswith("/login"))
page = anon.get("/login")
ok("login page is public", page.status_code == 200 and b"Log in" in page.data)
ok("login background picture is public", anon.get("/login-bg.jpg").status_code == 200)
ok("blueprint picture needs a login", anon.get("/blueprint-bg.jpg").status_code == 302)
ok("api says 401 when logged out", anon.get("/api/blocks").get_json() == {"error": "Please log in."})
ok("wrong password rejected", anon.post("/api/login", json={"username": "juan", "password": "bad"}).status_code == 401)
ok("unknown user rejected", anon.post("/api/login", json={"username": "zed", "password": "pw1"}).status_code == 401)
r = anon.post("/api/login", json={"username": "JUAN", "password": "pw1"})
ok("login works (username not case-sensitive)", r.status_code == 200 and r.get_json()["user"] == "juan")
ok("cookie is HttpOnly + SameSite", "HttpOnly" in r.headers["Set-Cookie"] and "SameSite=Lax" in r.headers["Set-Cookie"])
ok("logged-in browser opens the app", anon.get("/").status_code == 200)
ok("blueprint picture loads after login", anon.get("/blueprint-bg.jpg").status_code == 200)
ok("login page sends logged-in users to the app", anon.get("/login").status_code == 302)
anon.post("/api/logout")
ok("logout locks the app again", anon.get("/api/blocks").status_code == 401)

u = A.app.test_client()
u.post("/api/login", json={"username": "maria", "password": "pw2"})
ok("maria logged in", u.get("/api/blocks").status_code == 200)
A.USERS["maria"] = "newpw"
ok("changing a password signs that user out", u.get("/api/blocks").status_code == 401)
A.USERS["maria"] = "pw2"
ok("restoring it lets the old cookie work again", u.get("/api/blocks").status_code == 200)
del A.USERS["maria"]
ok("removing a user signs them out at once", u.get("/api/blocks").status_code == 401)
A.USERS["maria"] = "pw2"
print("ALL PASSED")
