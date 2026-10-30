"""Build attendance spreadsheets (.xlsx) from students and records."""
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
LABELS = [("Present", "P"), ("Absent", "A"), ("Late", "L"), ("Excused", "E")]
FILLS = {"P": "D9F2E3", "A": "F8D7D7", "L": "FBEBC8", "E": "E3DEF8"}


def _fill(ws, students, records):
    dates = sorted(records)
    first_date_col = 8  # A=ID, B=Name, C..F=counts, G=percent, H onward = dates
    headers = ["ID", "Name"] + [l for l, _ in LABELS] + ["Attendance %"] + dates
    ws.append(headers)

    head_fill = PatternFill("solid", fgColor="2F6FED")
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    last_col = get_column_letter(max(first_date_col, first_date_col + len(dates) - 1))
    first_col = get_column_letter(first_date_col)

    for r, s in enumerate(students, start=2):
        rng = f"${first_col}{r}:${last_col}{r}"
        row = [s["id"], s["name"]]
        row += [f'=COUNTIF({rng},"{code}")' for _, code in LABELS]
        row += [f"=IF(SUM(C{r}:F{r})=0,0,(C{r}+E{r}+F{r})/SUM(C{r}:F{r}))"]
        row += [records[d].get(s["id"], "") for d in dates]
        ws.append(row)
        ws.cell(r, 7).number_format = "0.00%"
        for c in range(1, len(headers) + 1):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT)
            if c >= first_date_col:
                cell.alignment = Alignment(horizontal="center")
                if cell.value in FILLS:
                    cell.fill = PatternFill("solid", fgColor=FILLS[cell.value])

    n = len(students) + 1
    ws.cell(n + 2, 1, "P = Present, A = Absent, L = Late, E = Excused. "
                      "Attendance % counts Present, Late and Excused as attended.").font = Font(name=FONT, italic=True)

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    for c in range(3, 8):
        ws.column_dimensions[get_column_letter(c)].width = 11
    for c in range(first_date_col, first_date_col + len(dates)):
        ws.column_dimensions[get_column_letter(c)].width = 12
    ws.freeze_panes = "C2"


def _fill_activities(ws, students, activities):
    acts = sorted(activities, key=lambda a: (a.get("date") or "9999-99-99", a["title"].lower()))
    fc = 5  # A=ID, B=Name, C=Done, D=Progress, E onward = one column per activity
    ws.append(["ID", "Name", "Done", "Progress"]
              + [f'{a["title"]}\n{a.get("date") or "no date"}' for a in acts])
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F6FED")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 45

    first = get_column_letter(fc)
    last = get_column_letter(max(fc, fc + len(acts) - 1))
    done_sets = [set(a.get("done", [])) for a in acts]
    for r, s in enumerate(students, start=2):
        hdr = f"${first}$1:${last}$1"
        ws.append([s["id"], s["name"], f'=COUNTIF(${first}{r}:${last}{r},"Done")',
                   f"=IF(COUNTA({hdr})=0,0,C{r}/COUNTA({hdr}))"]
                  + ["Done" if s["id"] in ds else "" for ds in done_sets])
        ws.cell(r, 4).number_format = "0.00%"
        for c in range(1, 5 + len(acts)):
            cell = ws.cell(r, c)
            cell.font = Font(name=FONT)
            if c >= fc:
                cell.alignment = Alignment(horizontal="center")
                if cell.value == "Done":
                    cell.fill = PatternFill("solid", fgColor=FILLS["P"])

    n = len(students) + 1
    if students and acts:
        ws.cell(n + 2, 2, "Students done per activity").font = Font(name=FONT, bold=True)
        for i in range(len(acts)):
            col = get_column_letter(fc + i)
            cell = ws.cell(n + 2, fc + i, f'=COUNTIF({col}2:{col}{n},"Done")')
            cell.font = Font(name=FONT, bold=True)
            cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 9
    ws.column_dimensions["D"].width = 11
    for i in range(len(acts)):
        ws.column_dimensions[get_column_letter(fc + i)].width = 18
    ws.freeze_panes = "C2"


def _title(name, used):
    base = (re.sub(r"[\[\]:*?/\\]", "", name).strip() or "Sheet")[:28]
    t, n = base, 2
    while t.lower() in used:
        t, n = f"{base} {n}", n + 1
    used.add(t.lower())
    return t


def build_workbook_multi(items):
    """items: [(sheet name, students, records[, activities]), ...] -> BytesIO.

    Each block gets an attendance sheet; blocks that have activities also get an
    "<name> Activities" sheet right after it.
    """
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for item in items:
        name, students, records = item[0], item[1], item[2]
        activities = item[3] if len(item) > 3 else []
        _fill(wb.create_sheet(_title(name, used)), students, records)
        if activities:
            _fill_activities(wb.create_sheet(_title(f"{name[:17]} Activities", used)), students, activities)
    if not wb.sheetnames:
        wb.create_sheet("Attendance")
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_workbook(students, records, name="Attendance", activities=None):
    return build_workbook_multi([(name, students, records, activities or [])])
